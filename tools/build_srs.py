from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('docs/Baganetic_SRS.docx')

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 89, 89)
LIGHT = 'F2F4F7'

def font(run, size=11, color=None, bold=None, italic=None, name='Calibri'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    margins = tcPr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar'); tcPr.append(margins)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = margins.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); margins.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    layout = tblPr.first_child_found_in('w:tblLayout')
    if layout is None: layout = OxmlElement('w:tblLayout'); tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    ind = tblPr.first_child_found_in('w:tblInd')
    if ind is None: ind = OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for gc, width in zip(grid.gridCol_lst, widths): gc.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width/1440)
            tcW = cell._tc.tcPr.tcW
            tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def para(doc, text='', style=None, before=0, after=6, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        a=p.add_run(bold_prefix); font(a, bold=True)
        b=p.add_run(text[len(bold_prefix):]); font(b)
    else:
        r=p.add_run(text); font(r)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.167
    p.add_run(text); return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.167
    p.add_run(text); return p

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT
    t.style='Table Grid'; set_table_widths(t,widths); set_repeat_table_header(t.rows[0])
    for cell, value in zip(t.rows[0].cells,headers):
        shade(cell,LIGHT); p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        r=p.add_run(value); font(r, size=9.5, color=INK, bold=True)
    for values in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,values):
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            r=p.add_run(str(value)); font(r,size=9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)

# Styles: standard_business_brief tokens.
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
for level,size,color,before,after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,DARK_BLUE,8,4)]:
    s=doc.styles[f'Heading {level}']; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); s.font.size=Pt(size); s.font.color.rgb=color; s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

# Header/footer.
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
r=hp.add_run('BAGANETIC  |  SOFTWARE REQUIREMENTS SPECIFICATION'); font(r,9,MUTED,bold=True)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=fp.add_run('Baganetic SRS  |  Version 1.0  |  01 September 2026'); font(r,8.5,MUTED)

# First-page memo masthead.
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(4)
r=p.add_run('SOFTWARE REQUIREMENTS SPECIFICATION'); font(r,23,INK,bold=True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(16)
r=p.add_run('Baganetic Mobile: Smart Route & Information Guide for Bagan Pagodas'); font(r,14,MUTED)
table(doc,['Document control','Value'],[
    ['Product name','Baganetic Mobile'],['Version','1.0'],['Status','Baseline specification for the current repository'],
    ['Prepared for','Special Project, I Semester 2025–2026'],['Prepared on','01 September 2026'],
    ['Primary platform','Flutter mobile application (Android-first; iOS, web, Windows, macOS, and Linux project targets present)']], [2520,6840])
para(doc,'Purpose of this document',before=8,after=4,bold_prefix='Purpose of this document')
para(doc,'This SRS defines the externally observable requirements for the implemented Baganetic Mobile application. It is based on the current source repository and distinguishes delivered functions from future enhancements mentioned in earlier project-planning material.',after=10)

heading(doc,'1. Introduction')
heading(doc,'1.1 Purpose',2)
para(doc,'Baganetic supports visitors exploring Bagan’s pagodas in environments where connectivity may be intermittent. The system combines an offline-first pagoda catalog, cached mapping, local route computation, trip planning, optional cloud-backed accounts, road reports, community content, and an authenticated AI travel assistant.')
heading(doc,'1.2 Scope',2)
para(doc,'The product is a Flutter application with a companion browser-based administration console. Core browsing, local catalog access, local trips, and locally available route data remain usable without Firebase. Authentication, road reports, cloud synchronization, and live AI replies depend on configured Firebase services and, for AI, a deployed Cloud Function.')
heading(doc,'1.3 Definitions and abbreviations',2)
table(doc,['Term','Meaning'],[
 ['A*','Graph pathfinding algorithm used for shortest-path calculation.'],['Firebase','Cloud platform used for authentication, Firestore, Storage, and Cloud Functions.'],['Firestore','Firebase’s document database used for synchronized application data.'],['Guest mode','Use of the app without signed-in Firebase credentials.'],['Offline graph','Bundled route-node and route-edge data used for local routing.'],['OSM','OpenStreetMap tile service used by the interactive map.'],['SRS','Software Requirements Specification.']], [2160,7200])
heading(doc,'1.4 Document conventions',2)
para(doc,'Each requirement is uniquely identified. “Shall” indicates a mandatory system behavior. “Should” indicates a quality target. Requirements marked “Conditional” apply only when their dependent service is configured and available.')

heading(doc,'2. Overall Description')
heading(doc,'2.1 Product perspective',2)
para(doc,'Baganetic is a standalone, offline-first tourism companion. The mobile client owns local SQLite data, cached map tiles, and bundled assets. Firebase augments the client with identity, shared data, and AI interaction. The administration console manages catalog, user, trip, and community content through the Firebase-backed environment.')
heading(doc,'2.2 User classes',2)
table(doc,['User class','Description','Key permissions'],[
 ['Guest visitor','User who has not signed in.','Browse pagodas, use map and local routes, create local trips, change language; cloud-gated actions prompt for sign-in.'],
 ['Registered visitor','Email/password authenticated visitor.','All guest functions plus synchronized identity-dependent functions such as road reports, community contributions, and AI chat history.'],
 ['Administrator','Authorized operator of the admin console.','Manage pagoda catalog, review users and trips, and moderate community posts.']], [1800,3000,4560])
heading(doc,'2.3 Operating environment',2)
bullet(doc,'Flutter and Dart application targeting Android; repository runners also support iOS, web, Windows, macOS, and Linux.')
bullet(doc,'SQLite/sqflite local storage; sqflite FFI is configured for Windows and Linux desktop targets.')
bullet(doc,'OpenStreetMap-compatible map tiles through flutter_map and a local tile cache.')
bullet(doc,'Firebase Authentication, Cloud Firestore, Cloud Storage, and Cloud Functions when cloud features are enabled.')
heading(doc,'2.4 Constraints and assumptions',2)
bullet(doc,'The first viewing of an uncached map area requires network access; previously cached tiles can be reused offline.')
bullet(doc,'Route availability and quality are limited to the bundled offline graph and supported multimodal route data.')
bullet(doc,'Live AI requires a signed-in user, Firebase configuration, a deployed generateAiReply function, and its server-side NVIDIA API key secret.')
bullet(doc,'The system does not promise live GPS navigation, live location sharing, hotels/restaurants discovery, emergency tooling, or group-trip chat in the current application baseline.')

heading(doc,'3. Functional Requirements')
heading(doc,'3.1 Application bootstrap, onboarding, and language',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-01','The system shall initialize the local experience even when Firebase initialization fails or times out.','High'],
 ['FR-02','The system shall show onboarding until the user marks it complete, then open the home experience on subsequent launches.','High'],
 ['FR-03','The system shall provide English and Myanmar locales and persist the user’s language preference.','High'],
 ['FR-04','The system shall present primary navigation for Home, Pagodas, Map/Trips, and Profile.','High']], [900,7020,1440])
heading(doc,'3.2 Pagoda catalog and detail information',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-05','The system shall seed and maintain a local catalog from bundled pagoda JSON data, including images and geographic coordinates.','High'],
 ['FR-06','The system shall allow users to browse, search, and filter pagodas by supported category.','High'],
 ['FR-07','The system shall display pagoda detail information, including description, historic and visitor guidance where available, images, and unique features.','High'],
 ['FR-08','The system shall allow a user to mark pagodas as favorites and retain favorites against the guest or signed-in owner identity.','High'],
 ['FR-09','The system shall show configured featured pagodas on the home screen.','Medium']], [900,7020,1440])
heading(doc,'3.3 Maps and offline routing',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-10','The system shall display an interactive map centered on the Bagan region with pagoda markers.','High'],
 ['FR-11','The system shall cache map tiles locally and perform background prefetch/cache initialization without blocking application startup.','High'],
 ['FR-12','The system shall generate routes using bundled offline graph data when route nodes and edges permit a path.','High'],
 ['FR-13','The system shall support available multimodal route data and calculate route distance using supported transport modes.','High'],
 ['FR-14','The system shall make route creation available from map and pagoda-detail flows.','High']], [900,7020,1440])
heading(doc,'3.4 Trip planning',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-15','The system shall let a guest or registered user create, edit, view, and delete locally stored trips.','High'],
 ['FR-16','The system shall store trip name, description, date, transport mode, timestamps, and ordered pagoda stops.','High'],
 ['FR-17','The system shall retain stop coordinates, notes, ordering, and distance-to-next information where calculated.','High'],
 ['FR-18','The system shall show a trip detail view with its stops and route-relevant information.','High']], [900,7020,1440])
heading(doc,'3.5 Identity and profile',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-19','The system shall permit guest-mode use for features that do not require a cloud identity.','High'],
 ['FR-20','Conditional: the system shall support Firebase email/password registration, sign-in, sign-out, and profile state when Firebase is configured.','High'],
 ['FR-21','The system shall present a profile area that reflects the current guest or signed-in state.','High'],
 ['FR-22','The system shall require sign-in before protected cloud features and explain the requirement to the user.','High']], [900,7020,1440])
heading(doc,'3.6 Road reports, feedback, and community',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-23','Conditional: signed-in users shall be able to create and view road reports associated with pagodas.','High'],
 ['FR-24','The system shall retain feedback records with pagoda, user identity, rating, comment, and timestamp fields.','Medium'],
 ['FR-25','The system shall provide a community feed of active tips and questions and a post-detail experience.','High'],
 ['FR-26','Conditional: signed-in users shall be able to create community posts and comments, and vote on posts subject to Firestore rules.','High'],
 ['FR-27','The system shall support administrative moderation of community-post status.','Medium']], [900,7020,1440])
heading(doc,'3.7 AI travel assistant',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-28','The system shall expose an AI assistant entry point from supported context flows.','Medium'],
 ['FR-29','Conditional: the system shall require an authenticated Firebase user before requesting a live AI reply.','High'],
 ['FR-30','Conditional: the system shall send only the required user message and supported pagoda, trip, or road context to the generateAiReply Cloud Function.','High'],
 ['FR-31','The system shall retain local AI message history and synchronize eligible history when cloud services are available.','Medium'],
 ['FR-32','The mobile application shall not embed the external AI provider key in source, configuration defines, or a distributable build.','High']], [900,7020,1440])
heading(doc,'3.8 Administration console',2)
table(doc,['ID','Requirement','Priority'],[
 ['FR-33','The administration console shall let authorized operators list, search, create, edit, archive, and feature pagoda catalog entries.','High'],
 ['FR-34','The administration console shall support reviewing registered users, saved trips, and community-post activity.','Medium'],
 ['FR-35','The administration console shall provide moderation actions for community content and a confirmation step for destructive actions.','High']], [900,7020,1440])

heading(doc,'4. External Interface Requirements')
heading(doc,'4.1 User interface',2)
bullet(doc,'Mobile UI shall use responsive Flutter layouts and provide navigation appropriate to the active form factor.')
bullet(doc,'Screens shall use the defined application theme and provide readable English and Myanmar text using the bundled font assets.')
bullet(doc,'Map interactions shall communicate selected places, route paths, and navigation affordances without requiring a separate map application.')
heading(doc,'4.2 Software and service interfaces',2)
table(doc,['Interface','Purpose','Failure behavior'],[
 ['SQLite / sqflite','Local catalog, trips, favorites, feedback, community cache, and local AI messages.','Local database errors are logged; the app should remain recoverable where possible.'],
 ['OpenStreetMap tile endpoint','Map imagery and tile caching.','Show available cached imagery or normal map loading/error state.'],
 ['Firebase Authentication','Email/password identity.','Guest mode remains available.'],
 ['Cloud Firestore','Cloud-backed reports, community data, and optional synchronized history.','Cloud-dependent actions report an actionable unavailable/sign-in state.'],
 ['Cloud Function generateAiReply','Authenticated AI replies.','Show request failure without exposing secrets or raw provider credentials.'],
 ['Firebase Storage','Cloud-hosted media where configured.','Keep local/bundled content usable where available.']], [2400,3300,3660])
heading(doc,'4.3 Data interfaces',2)
para(doc,'Bundled inputs include sample pagoda JSON, offline route graph JSON, pagoda route graph JSON, multimodal routes JSON, and pagoda image assets. The offline graph may optionally be refreshed from a configured hosted source; the baseline configuration leaves that source empty, so routing graph updates are local-only.')

heading(doc,'5. Data Requirements')
table(doc,['Data entity','Minimum data handled','Primary store'],[
 ['Pagoda','ID, name, category, coordinates, description, images, historic/visitor fields, featured status.','SQLite; optional Firestore catalog/admin source'],
 ['Trip and stop','Owner, name, date, transport mode, ordered stops, coordinates, notes, distances.','SQLite'],
 ['Favorite','Owner identity, pagoda ID, timestamp.','SQLite'],
 ['Road report','Pagoda association, signed-in user data, report details, timestamp.','Firestore when configured'],
 ['Community post/comment/vote','Author, type, title/body, status, tags, timestamps, relationships.','SQLite cache and Firestore when configured'],
 ['AI message','Owner, role, text, contextual IDs, request ID, timestamp.','SQLite; optional Firestore synchronization']], [1800,5040,2520])
heading(doc,'5.1 Data validation and integrity',2)
bullet(doc,'Pagoda IDs shall be unique. Trip stop ordering shall be retained as an explicit index.')
bullet(doc,'Ratings, coordinates, owner identifiers, and required display fields shall be validated before persistence by the corresponding model/service flow.')
bullet(doc,'Database migrations shall be applied in version order, and the application shall attempt schema self-healing for critical prior schema drift.')
bullet(doc,'Cloud access shall be constrained by deployed Firebase Authentication, Firestore, and Storage rules.')

heading(doc,'6. Non-functional Requirements')
table(doc,['ID','Requirement','Verification approach'],[
 ['NFR-01','Offline resilience: core catalog, local trips, bundled routes, and previously cached map tiles shall remain available without active network connectivity.','Launch and use the app with network disabled after initial local setup/cache.'],
 ['NFR-02','Startup resilience: a Firebase initialization failure shall not prevent the local experience from becoming available.','Run with invalid/unavailable Firebase configuration.'],
 ['NFR-03','Performance: local catalog search and route calculation should complete responsively for the packaged Bagan data set on supported devices.','Measure interactions on a representative Android device/emulator.'],
 ['NFR-04','Security: external AI credentials shall reside only in Firebase Functions Secret Manager, not the mobile binary.','Inspect app source/build configuration and deployed function secret setup.'],
 ['NFR-05','Privacy: identity-dependent data shall be associated with the authenticated user and protected by Firebase rules.','Review rules and perform authorized/unauthorized access tests.'],
 ['NFR-06','Accessibility: the UI should use readable contrast, semantic labels for meaningful images, touch-friendly controls, and localized text.','Manual accessibility review on supported platforms.'],
 ['NFR-07','Maintainability: Flutter analysis and automated tests shall be runnable from the repository.','Run flutter analyze and flutter test.'],
 ['NFR-08','Reliability: background tile-cache and graph-sync failures shall be caught and logged rather than terminating the app.','Simulate service/caching failure and inspect logs.']], [900,5040,3420])

heading(doc,'7. Security and Privacy Requirements')
numbered(doc,'The client shall use Firebase Authentication state to gate protected cloud functions.')
numbered(doc,'The AI provider secret shall be configured server-side as NVIDIA_API_KEY through Firebase Functions Secret Manager.')
numbered(doc,'The application shall use Firestore and Storage rules appropriate to the user/administrator access model before production deployment.')
numbered(doc,'The system shall not log or display credentials, tokens, or raw server error data to end users.')
numbered(doc,'Administrators shall confirm destructive console actions before data deletion or other irreversible operations.')

heading(doc,'8. Acceptance Criteria')
table(doc,['Area','Acceptance criterion'],[
 ['Local-first experience','With Firebase unavailable, a new or returning user can complete onboarding, browse the bundled pagoda catalog, view details, use local trip functions, and access bundled routing data.'],
 ['Mapping and routes','The map displays pagoda locations; after tiles are cached, cached map areas remain visible offline; a supported source/destination pair produces a route.'],
 ['Trip planning','A user can create a trip, add ordered pagoda stops, save it, reopen it, and view the saved data.'],
 ['Authentication','With Firebase configured, a user can register/sign in/out; protected actions appropriately require sign-in.'],
 ['Community and reports','With Firebase configured and authorized, a signed-in user can submit a road report and create/engage with community content.'],
 ['AI','With all AI prerequisites configured, a signed-in user receives a reply from the Cloud Function; an unauthenticated user cannot make a live request.'],
 ['Quality gate','flutter analyze and flutter test complete successfully in the configured development environment.']], [2160,7200])

heading(doc,'9. Out of Scope for This Baseline')
para(doc,'The following concepts appear in earlier planning material but are not treated as delivered requirements of the current repository: hotel/restaurant/ATM discovery, an optimized interest-based itinerary engine, trip rooms and group chat, live location sharing, an emergency toolkit, live GPS turn-by-turn navigation, and AR guidance. They may be documented in a future release SRS after implementation decisions and acceptance criteria are defined.')

heading(doc,'10. Traceability and Validation')
para(doc,'The repository’s current automated checks are flutter analyze and flutter test. Tests cover onboarding preference, pagoda localization and home-feature behavior, routing/pathfinding, ferry-distance calculations, multimodal route services, and AI chatbot-service behavior. Release validation should also exercise Firebase-dependent scenarios using a test project and rules aligned with the intended deployment.')
heading(doc,'11. Revision History')
table(doc,['Version','Date','Description'],[['1.0','01 September 2026','Initial SRS baseline prepared from the Baganetic Mobile repository.']], [1440,2160,5760])

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
